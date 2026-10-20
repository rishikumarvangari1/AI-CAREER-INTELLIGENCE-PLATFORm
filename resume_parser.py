
import csv
import io
import json
import os
import re
import zipfile
from html import unescape
from pathlib import Path
from xml.etree import ElementTree

import fitz
import pytesseract
from PIL import Image, ImageOps, ImageEnhance
from docx import Document
from pptx import Presentation
from openpyxl import load_workbook
from striprtf.striprtf import rtf_to_text
from bs4 import BeautifulSoup


SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".txt", ".rtf",
    ".pptx", ".xlsx", ".xlsm", ".csv",
    ".tsv", ".json", ".xml", ".html", ".htm",
    ".md", ".log", ".yaml", ".yml",
    ".png", ".jpg", ".jpeg", ".webp",
    ".bmp", ".tif", ".tiff",
    ".odt", ".ods", ".odp"
}

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_CHARACTERS = 1_000_000


class ResumeParsingError(Exception):
    """Raised when a resume cannot be extracted safely."""


if os.name == "nt":
    TESSERACT_PATH = r"C:\Program Files\Tesseract-OCR\tesseract.exe"

    if os.path.exists(TESSERACT_PATH):
        pytesseract.pytesseract.tesseract_cmd = TESSERACT_PATH


def clean_text(text):
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n[ \t]+\n", "\n\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def check_file(filepath):
    path = Path(filepath)

    if not path.is_file():
        raise ResumeParsingError(
            "The uploaded file could not be found."
        )

    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ResumeParsingError(
            f"Unsupported format: {path.suffix or 'unknown'}. "
            "Upload a supported document or image format."
        )

    if path.stat().st_size > MAX_FILE_BYTES:
        raise ResumeParsingError(
            "The file exceeds the 10 MB limit."
        )

    if path.stat().st_size == 0:
        raise ResumeParsingError(
            "The uploaded file is empty."
        )

    return path


def text_quality(text):
    if not text or not text.strip():
        return 0

    words = text.split()

    readable = sum(
        character.isalnum()
        or character.isspace()
        or character in ".,:;()@&+-/'"
        for character in text
    )

    quality = readable / max(len(text), 1)

    useful_words = sum(
        any(character.isalpha() for character in word)
        for word in words
    )

    return quality * 100 + min(useful_words, 100)


def extract_pdf(path):
    pages_text = []

    try:
        with fitz.open(path) as pdf:
            if pdf.is_encrypted:
                raise ResumeParsingError(
                    "This PDF is password-protected. "
                    "Upload an unlocked copy."
                )

            for page in pdf:
                original_text = page.get_text("text").strip()
                extracted_text = original_text

                should_use_ocr = (
                    len(original_text) < 100
                    or text_quality(original_text) < 65
                )

                if should_use_ocr:
                    try:
                        pix = page.get_pixmap(
                            matrix=fitz.Matrix(3, 3),
                            alpha=False
                        )

                        image = Image.open(
                            io.BytesIO(pix.tobytes("png"))
                        )

                        image = ImageOps.grayscale(image)
                        image = ImageEnhance.Contrast(
                            image
                        ).enhance(2)

                        ocr_text = pytesseract.image_to_string(
                            image,
                            lang="eng",
                            config="--psm 3"
                        ).strip()

                        if text_quality(ocr_text) > text_quality(
                            original_text
                        ):
                            extracted_text = ocr_text

                    except pytesseract.TesseractNotFoundError as exc:
                        if not original_text:
                            raise ResumeParsingError(
                                "Scanned PDF detected, but Tesseract "
                                "OCR is not configured. Install "
                                "Tesseract OCR and verify its path."
                            ) from exc

                    except pytesseract.TesseractError as exc:
                        if not original_text:
                            raise ResumeParsingError(
                                "OCR could not process this scanned PDF. "
                                "Check that Tesseract and its English "
                                "language data are installed."
                            ) from exc

                    except Exception as exc:
                        if not original_text:
                            raise ResumeParsingError(
                                "Could not extract text from this "
                                "scanned PDF. Check the PDF and OCR setup."
                            ) from exc

                pages_text.append(extracted_text)

    except ResumeParsingError:
        raise

    except Exception as exc:
        raise ResumeParsingError(
            f"Unable to read PDF: {exc}"
        ) from exc

    return "\n".join(pages_text).strip()


def extract_docx(path):
    try:
        document = Document(path)
        parts = [p.text for p in document.paragraphs]

        for table in document.tables:
            for row in table.rows:
                parts.append(
                    " | ".join(cell.text for cell in row.cells)
                )

        for section in document.sections:
            parts.extend(
                p.text for p in section.header.paragraphs
            )
            parts.extend(
                p.text for p in section.footer.paragraphs
            )

        return "\n".join(parts)

    except Exception as exc:
        raise ResumeParsingError(
            "Could not read this Word document."
        ) from exc


def extract_pptx(path):
    try:
        presentation = Presentation(path)
        parts = []

        for slide in presentation.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False):
                    parts.append(shape.text)

                if getattr(shape, "has_table", False):
                    for row in shape.table.rows:
                        parts.append(
                            " | ".join(
                                cell.text for cell in row.cells
                            )
                        )

        return "\n".join(parts)

    except Exception as exc:
        raise ResumeParsingError(
            "Could not read this PowerPoint presentation."
        ) from exc


def extract_spreadsheet(path):
    try:
        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True
        )

        parts = []

        try:
            for sheet in workbook.worksheets:
                parts.append(f"Sheet: {sheet.title}")

                for row in sheet.iter_rows(values_only=True):
                    values = [
                        str(value)
                        for value in row
                        if value is not None
                    ]

                    if values:
                        parts.append(" | ".join(values))

        finally:
            workbook.close()

        return "\n".join(parts)

    except Exception as exc:
        raise ResumeParsingError(
            "Could not read this spreadsheet."
        ) from exc


def extract_image(path):
    try:
        with Image.open(path) as original:
            image = ImageOps.exif_transpose(original)
            image.load()

            if image.width * image.height > 40_000_000:
                raise ResumeParsingError(
                    "The image dimensions are too large for OCR."
                )

            image = image.convert("RGB")
            image = ImageOps.grayscale(image)
            image = ImageEnhance.Contrast(image).enhance(2)

            return pytesseract.image_to_string(
                image,
                lang="eng",
                config="--psm 3"
            )

    except ResumeParsingError:
        raise

    except pytesseract.TesseractNotFoundError as exc:
        raise ResumeParsingError(
            "Tesseract OCR is not installed or configured. "
            "Install Tesseract and verify its executable path."
        ) from exc

    except Exception as exc:
        raise ResumeParsingError(
            "Image text extraction failed. Check that OCR is "
            "installed and that the image is valid."
        ) from exc


def extract_text_file(path):
    raw = path.read_bytes()

    for encoding in (
        "utf-8-sig",
        "utf-8",
        "cp1252",
        "latin-1"
    ):
        try:
            return raw.decode(encoding)

        except UnicodeDecodeError:
            continue

    raise ResumeParsingError(
        "The text file encoding is unsupported."
    )


def extract_html(path):
    source = extract_text_file(path)
    soup = BeautifulSoup(source, "html.parser")

    for element in soup(
        ["script", "style", "noscript"]
    ):
        element.decompose()

    return soup.get_text(separator="\n")


def extract_rtf(path):
    try:
        return rtf_to_text(extract_text_file(path))

    except Exception as exc:
        raise ResumeParsingError(
            "Could not extract text from this RTF document."
        ) from exc


def extract_delimited(path, delimiter):
    text = extract_text_file(path)
    rows = csv.reader(
        io.StringIO(text),
        delimiter=delimiter
    )

    return "\n".join(
        " | ".join(cell for cell in row)
        for row in rows
    )


def extract_json(path):
    try:
        data = json.loads(extract_text_file(path))

        return json.dumps(
            data,
            ensure_ascii=False,
            indent=2
        )

    except Exception as exc:
        raise ResumeParsingError(
            "Invalid JSON document."
        ) from exc


def extract_xml(path):
    try:
        source = extract_text_file(path)
        root = ElementTree.fromstring(source)

        return "\n".join(
            text.strip()
            for text in root.itertext()
            if text.strip()
        )

    except Exception as exc:
        raise ResumeParsingError(
            "Could not read this XML document."
        ) from exc


def extract_open_document(path):
    """
    Extract text from ODT, ODS and ODP files.
    These formats are ZIP containers with XML content.
    """
    try:
        with zipfile.ZipFile(path) as archive:
            if "content.xml" not in archive.namelist():
                raise ResumeParsingError(
                    "The OpenDocument file has no content.xml."
                )

            xml_data = archive.read("content.xml")
            root = ElementTree.fromstring(xml_data)

            parts = []

            for element in root.iter():
                if (
                    element.tag.endswith("}p")
                    or element.tag.endswith("}h")
                ):
                    value = "".join(element.itertext()).strip()

                    if value:
                        parts.append(value)

            return "\n".join(parts)

    except ResumeParsingError:
        raise

    except Exception as exc:
        raise ResumeParsingError(
            "Could not read this OpenDocument file."
        ) from exc


def parse_resume(filepath):
    path = check_file(filepath)
    ext = path.suffix.lower()

    if ext == ".pdf":
        text = extract_pdf(path)

    elif ext == ".docx":
        text = extract_docx(path)

    elif ext == ".pptx":
        text = extract_pptx(path)

    elif ext in {".xlsx", ".xlsm"}:
        text = extract_spreadsheet(path)

    elif ext in {
        ".png", ".jpg", ".jpeg", ".webp",
        ".bmp", ".tif", ".tiff"
    }:
        text = extract_image(path)

    elif ext in {
        ".txt", ".md", ".log", ".yaml", ".yml"
    }:
        text = extract_text_file(path)

    elif ext == ".xml":
        text = extract_xml(path)

    elif ext in {".html", ".htm"}:
        text = extract_html(path)

    elif ext == ".rtf":
        text = extract_rtf(path)

    elif ext in {".csv", ".tsv"}:
        text = extract_delimited(
            path,
            "\t" if ext == ".tsv" else ","
        )

    elif ext == ".json":
        text = extract_json(path)

    elif ext in {".odt", ".ods", ".odp"}:
        text = extract_open_document(path)

    else:
        raise ResumeParsingError(
            "This file format is not supported."
        )

    text = clean_text(text)

    if not text:
        raise ResumeParsingError(
            "No readable text was found. Try a clearer document "
            "or an OCR-enabled version."
        )

    if len(text) > MAX_EXTRACTED_CHARACTERS:
        text = text[:MAX_EXTRACTED_CHARACTERS]

    return text
